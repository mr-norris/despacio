"""Saves every tab of the Mirror spreadsheet as a CSV file in data/source/.

Run daily by .github/workflows/backup.yml, which commits any changes, so the repo
holds a dated history of the community sheet. Reads the Mirror's "Publish to web"
link (entire document, .xlsx), whose ID is mirrorPubId in config.json.
"""
import csv, io, json, re, sys, urllib.request
from pathlib import Path

import openpyxl

pub_id = json.loads(Path('config.json').read_text()).get('mirrorPubId', '').strip()
if not pub_id:
    print('No mirrorPubId in config.json yet; nothing to back up.')
    sys.exit(0)

url = f'https://docs.google.com/spreadsheets/d/e/{pub_id}/pub?output=xlsx'
data = urllib.request.urlopen(url, timeout=120).read()
if not data.startswith(b'PK'):  # an .xlsx is a zip file; anything else is an error page
    sys.exit(f'Expected a spreadsheet from {url}, got something else. Is the Mirror still published as .xlsx?')

book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
if not book.sheetnames:
    sys.exit('The published Mirror has no tabs; keeping the previous backup.')

out = Path('data/source')
out.mkdir(parents=True, exist_ok=True)
slug = lambda name: re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-') or 'tab'

index, written = [], set()
for ws in book.worksheets:
    rows = [['' if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()  # trailing blank rows
    name = slug(ws.title)
    while f'{name}.csv' in written:
        name += '-2'
    written.add(f'{name}.csv')
    buf = io.StringIO()
    csv.writer(buf, lineterminator='\n').writerows(rows)
    (out / f'{name}.csv').write_text(buf.getvalue(), encoding='utf-8')
    index.append([ws.title, f'{name}.csv', sum(1 for r in rows if any(c.strip() for c in r))])

# Tabs removed from the Mirror are removed here too; git history still has them.
for f in out.glob('*.csv'):
    if f.name not in written and f.name != '_tabs.csv':
        f.unlink()

buf = io.StringIO()
csv.writer(buf, lineterminator='\n').writerows([['Tab', 'File', 'Rows']] + index)
(out / '_tabs.csv').write_text(buf.getvalue(), encoding='utf-8')
print(f'Saved {len(index)} tabs to {out}/.')
